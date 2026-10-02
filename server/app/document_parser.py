from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}


class DocumentParseError(ValueError):
    pass


def parse_document(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(
            f"Unsupported file type {extension or '(none)'}. "
            f"Allowed types: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    try:
        if extension in {".txt", ".md"}:
            return content.decode("utf-8-sig").strip()
        if extension == ".docx":
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append("\t".join(cell.text for cell in row.cells))
            return "\n".join(part for part in parts if part.strip()).strip()
        reader = PdfReader(BytesIO(content))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as exc:
        raise DocumentParseError(f"Could not parse {filename}: {exc}") from exc
