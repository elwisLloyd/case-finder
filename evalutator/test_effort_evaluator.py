import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from docx import Document
from pydantic import ValidationError

from effort_evaluator import (COLUMNS, MENU_COLUMNS, Mark, Marks, Summary, Addition,
                              evaluate, load_menu, load_spec, save_results, save_summary, statistics)


class EvaluatorTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.menu = pd.DataFrame([
            ['Каналы', 'Веб-чат', 'Нет', 'Виджет', 'Кастомный', 'Омниканал'],
            ['Голос', 'STT', 'Нет', 'API', 'On-prem', 'Обучение'],
        ], columns=MENU_COLUMNS)

    def test_docx_tables_and_containment(self):
        doc = Document()
        doc.add_paragraph('Начало ТЗ')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Требование из таблицы'
        doc.add_paragraph('Конец ТЗ')
        path = self.root / 'spec.DOCX'
        doc.save(path)
        text = load_spec(path, self.root)
        self.assertLess(text.index('Начало'), text.index('Требование'))
        self.assertLess(text.index('Требование'), text.index('Конец'))
        with self.assertRaises(ValueError):
            load_spec(path, self.root / 'TZ')

    def test_menu_and_csv(self):
        path = self.root / 'menu.csv'
        self.menu.to_csv(path, index=False, encoding='utf-8-sig')
        self.assertEqual(len(load_menu(path)), 2)
        pd.concat([self.menu, self.menu]).to_csv(path, index=False)
        with self.assertRaises(ValueError):
            load_menu(path)

    def test_id_order_statistics_and_save(self):
        response = Marks(items=[
            Mark(item_id=1, rating='не определено', confidence=.8, comment='Не указан голос.'),
            Mark(item_id=0, rating='small', confidence=.9, comment='Требуется виджет.'),
        ])
        with patch('effort_evaluator.ask', return_value=response):
            result = evaluate(None, 'test', self.menu, 'ТЗ')
        self.assertEqual(result['Оценка'].tolist(), ['small', 'не определено'])
        self.assertEqual(statistics(result)['Успешно проставлено'], 1)
        path = save_results(result, self.root, 'test')
        self.assertEqual(pd.read_csv(path).columns.tolist(), COLUMNS)
        summary = Summary(conclusion='Нужны уточнения.', questions=['Нужен ли голос?'], additions=[
            Addition(block='Новый блок', component='Новый пункт', reason='Требование ТЗ',
                     absent='Не требуется', small='Простой', medium='Средний', large='Сложный')])
        additions, report = save_summary(summary, result, self.root, 'test')
        self.assertEqual(additions.loc[0, 'Large'], 'Сложный')
        self.assertIn('Не определены: 1', report.read_text())

    def test_bad_ids_and_confidence_rejected(self):
        mark = Mark(item_id=0, rating='small', confidence=.9, comment='Виджет')
        for items in [[mark], [mark, mark], [mark, mark.model_copy(update={'item_id': 5})]]:
            with patch('effort_evaluator.ask', return_value=Marks(items=items)):
                with self.assertRaises(ValueError):
                    evaluate(None, 'test', self.menu, 'ТЗ')
        with self.assertRaises(ValidationError):
            Mark(item_id=0, rating='small', confidence=1.1, comment='Ошибка')
        with self.assertRaises(ValueError):
            evaluate(None, 'test', self.menu, 'Длинный текст', max_spec_chars=2)


if __name__ == '__main__':
    unittest.main()
